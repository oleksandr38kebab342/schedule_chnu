from datetime import time
from io import BytesIO

from openpyxl import load_workbook

from app.models import Parity
from app.services.excel_export import (
    Lesson,
    PairSlot,
    build_grid,
    generate_workbook,
)

PAIRS = [
    PairSlot(1, time(8, 20), time(9, 40)),
    PairSlot(2, time(9, 50), time(11, 10)),
]


def test_weekdays_always_present_even_without_lessons():
    grid = build_grid([], PAIRS)
    assert grid.days == [1, 2, 3, 4, 5]


def test_saturday_added_only_when_lesson_exists():
    grid = build_grid([Lesson(6, 1, Parity.ODD, "Субота")], PAIRS)
    assert 6 in grid.days
    assert 7 not in grid.days


def test_both_parity_fills_odd_and_even():
    grid = build_grid([Lesson(1, 1, Parity.BOTH, "Лекція")], PAIRS)
    assert grid.cell(1, 1, Parity.ODD) == ["Лекція"]
    assert grid.cell(1, 1, Parity.EVEN) == ["Лекція"]


def test_odd_only_does_not_fill_even():
    grid = build_grid([Lesson(1, 1, Parity.ODD, "Тільки непарний")], PAIRS)
    assert grid.cell(1, 1, Parity.ODD) == ["Тільки непарний"]
    assert grid.cell(1, 1, Parity.EVEN) == []


def test_overlaps_collected_not_overwritten():
    # DoD: накладки виводяться списком, а не перезаписом клітинки.
    lessons = [
        Lesson(1, 1, Parity.ODD, "Заняття A"),
        Lesson(1, 1, Parity.ODD, "Заняття B"),
    ]
    grid = build_grid(lessons, PAIRS)
    assert grid.cell(1, 1, Parity.ODD) == ["Заняття A", "Заняття B"]


def test_nonstandard_time_slot_survives_empty_label():
    # Пара без заданого часу: слот лишається, підпис порожній (час не губиться на інших).
    grid = build_grid([], [PairSlot(3, None, None)])
    assert grid.pairs[0].number == 3
    assert grid.pairs[0].time_label() == ""


def test_workbook_is_valid_xlsx_with_header():
    grid = build_grid([Lesson(1, 1, Parity.ODD, "Математика\nІваненко\nА.203")], PAIRS)
    content = generate_workbook([("121", grid)])
    assert content[:2] == b"PK"  # xlsx = zip

    wb = load_workbook(BytesIO(content))
    ws = wb["121"]
    assert ws["A2"].value == "Пара"
    assert ws["B2"].value == "Час"
    assert ws["C2"].value == "Понеділок"
    assert ws.freeze_panes == "C3"
    assert ws.page_setup.orientation == "landscape"


def test_empty_sheets_rejected():
    import pytest

    with pytest.raises(ValueError):
        generate_workbook([])


def test_multiple_group_sheets():
    g1 = build_grid([Lesson(1, 1, Parity.ODD, "A")], PAIRS)
    g2 = build_grid([Lesson(2, 1, Parity.ODD, "B")], PAIRS)
    content = generate_workbook([("121", g1), ("122", g2)])
    wb = load_workbook(BytesIO(content))
    assert wb.sheetnames == ["121", "122"]
